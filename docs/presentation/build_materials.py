#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate editable contest materials and matching layout previews.

Dependencies: python-pptx, python-docx, reportlab, Pillow.
All figures are editable shapes; no fabricated board screenshots are used.
"""
import json
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DColor
from docx.oxml import OxmlElement as DX
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4

OUT = Path(__file__).resolve().parent
FONT = Path('/mnt/c/Windows/Fonts/msyh.ttc')
BOLD = Path('/mnt/c/Windows/Fonts/msyhbd.ttc')
PDF_FONT = Path('/mnt/c/Windows/Fonts/simhei.ttf')
BG, INK, MUTED = 'F7F9FA', '172B32', '50636C'
TEAL, RED, LINE, WHITE = '007F78', 'B64637', 'D9E2E5', 'FFFFFF'
W, H = 13.333333, 7.5
slides = []


def new(title, kicker, subtitle='', seconds=20, notes=''):
    page = {'title': title, 'seconds': seconds, 'notes': notes, 'items': []}
    slides.append(page)
    box(0, 0, W, H, BG)
    box(.48, .40, .10, .30, TEAL)
    text(.72, .40, 11.8, .3, kicker, 11, TEAL, True)
    text(.7, .98, 11.9, .68, title, 29, INK, True)
    if subtitle:
        text(.72, 1.78, 11.85, .58, subtitle, 15, MUTED)
    box(.7, 6.98, 11.9, .014, LINE)
    text(.72, 7.12, 10.9, .18, 'OPENVELA / ESP32-P4   ·   队伍 124 神一样的老师   ·   2026.09.16', 9, MUTED)
    text(12.0, 7.10, .6, .2, f'{len(slides):02d}', 10, TEAL, True)
    return page


def box(x, y, w, h, fill, stroke=None):
    slides[-1]['items'].append(dict(kind='box', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke))


def text(x, y, w, h, value, size=18, color=INK, bold=False):
    slides[-1]['items'].append(dict(kind='text', x=x, y=y, w=w, h=h, value=value, size=size, color=color, bold=bold))


def block(x, y, width, label, body, accent=TEAL):
    box(x, y, .055, .37, accent)
    text(x+.18, y, width-.18, .42, label, 19, accent, True)
    text(x+.18, y+.62, width-.18, 1.15, body, 16, INK)


def placeholder(x, y, w, h, number, title, hint):
    box(x, y, w, h, 'EDF2F3', 'B3C7CC')
    text(x+.25, y+.26, w-.5, .30, f'截图位 {number} / 待补实机画面', 12, TEAL, True)
    text(x+.25, y+.90, w-.5, .85, title, 22, INK, True)
    text(x+.25, y+1.96, w-.5, h-2.12, hint, 13, MUTED)


new('ESP32-P4 上的 openvela AI 设备诊断助手', '作品介绍 / 新硬件适配 + AI 应用',
    '让模型读取真实设备状态，让一次问答可靠回到命令行。', 18,
    '我们没有把目标停在系统启动或接入聊天接口，而是面向无屏开发板，把平台适配、系统可靠性与真实状态工具连成一个可演示的闭环。项目基于上游开源成果，不宣称首个 ESP32-P4 适配。')
block(.75, 2.85, 6.2, '从“能启动”走向“可诊断”', 'USB 控制台 · 有线联网 · 板端工具\nopenvela Agent · 云端模型 · NSH 返回')
text(.95, 5.32, 5.8, .8, '本项目贡献：板级集成、运行时修正、\n设备工具与可复用调试方法。', 18, TEAL, True)
placeholder(7.7, 2.55, 4.85, 3.95, '01', '开发板实物与连接', '补图：ESP32-P4、USB 与网线。\n不用放 API Key 或配置截图。')

new('相同芯片，交付深度可以不同', '01 / 场景与差异化',
    '面向开发者的无屏设备诊断，不是通用聊天终端。', 20,
    '其他开发者也在做 ESP32-P4 适配，因此不能仅以芯片型号体现差异。我们的增量是让设备状态可被模型读取，并解决实际运行中的栈、堆、TLS、时间和任务生命周期问题。这里比较的是交付层次，不评价其他参赛者。')
block(.8, 2.7, 3.8, '基础：平台能运行', '启动 NSH\nUSB Console\n以太网与文件系统')
block(4.9, 2.7, 3.6, '进阶：Agent 能执行', '查询真实堆与运行时间\n读取 eth0 状态\n通过 JSON 返回给模型')
block(8.9, 2.7, 3.65, '重点：过程可复现', '异常寄存器与堆链表证据\n同步 ask 返回 NSH\n构建、配置与验收闭环')
box(.8, 5.45, 11.7, .84, TEAL)
text(1.02, 5.66, 11.2, .43, '差异化来自“真实工具 + 运行可靠性 + 调试证据”，不是重复声明移植成功。', 19, WHITE, True)

new('保留上游归属，明确本项目增量', '02 / 原创边界',
    '基于开源平台做工程集成与针对性改进，不将上游驱动或 Agent 引擎算作原创。', 20,
    'NuttX 和 Espressif 提供芯片及驱动基础，openvela 提供 Agent、工具注册与模型调用框架。本项目工作是将它们在选定板卡上连通，并针对实机问题修正。正式提交还需要固定基线版本并附相关 PR 或提交记录。')
text(.9, 2.58, 5.4, .4, '复用的基础', 22, MUTED, True)
text(.9, 3.24, 5.35, 2.2, 'NuttX / Espressif 的 ESP32-P4 基础\nopenvela ai_agent 与 ReAct 工具循环\ncJSON、mbedTLS、NTP 等组件\n已有板卡适配与社区工作', 18)
box(6.35, 2.65, .018, 3.55, LINE)
text(6.75, 2.58, 5.4, .4, '本项目实现 / 修正', 22, TEAL, True)
text(6.75, 3.24, 5.5, 2.7, 'USB + EMAC + 存储的目标配置\ndevice_info / network_status 工具\nCLI 同步 ask 与资源生命周期\n栈 / 堆 / USB / TLS / 时间适配\n实机定位与验收流程', 18)

new('设备端执行，云端推理', '03 / 系统架构',
    '模型通过工具获取观测值，而不是仅根据提示词猜测设备状态。', 23,
    '用户从 USB NSH 发起 ask。ESP32-P4 上的 openvela Agent 负责模型请求和工具调度，华为 MaaS 承担推理。本地工具读取系统 API 并输出 JSON，模型解释结果。问答结束后返回原 NSH。网络状态工具只报告接口配置，不等于主动测通外网。')
for x, title, body in [(0.8, 'USB / NSH', '用户输入\nai_agent ask'), (3.85, 'openvela Agent', '同步请求\nReAct 工具调度'), (8.6, '华为 MaaS', 'HTTPS / TLS 1.2\n云端模型推理')]:
    box(x, 2.65, 2.85 if x != 3.85 else 3.25, 1.30, WHITE, LINE)
    text(x+.18, 2.81, 2.5, .35, title, 19, TEAL, True)
    text(x+.18, 3.26, 2.6, .6, body, 14)
text(3.66, 3.10, .19, .3, '→', 10, TEAL, True)
text(7.30, 3.05, 1.1, .45, '← →', 23, TEAL, True)
text(5.2, 4.53, 1.0, .30, '↑ ↓', 16, TEAL, True)
text(3.0, 4.2, 8.4, .4, '请求 / 工具调用 / JSON 结果 / 自然语言解释', 16, MUTED)
box(3.85, 4.95, 7.6, 1.15, TEAL)
text(4.1, 5.14, 7.05, .34, 'device_info       network_status', 21, WHITE, True)
text(4.1, 5.64, 7.05, .28, 'uname / mallinfo / CLOCK_MONOTONIC / netlib', 15, WHITE)
text(.9, 5.13, 2.6, .75, 'SmartFS /data\n配置与会话数据', 17)

new('两项工具，把语言连接到设备事实', '04 / 本项目工具扩展',
    '不依赖额外传感器，以板卡已有的系统资源和以太网作为可验证输入。', 22,
    'device_info 获取内核信息、运行时间、堆总量、使用量和最大空闲块。network_status 获取网口注册和运行状态、IP、掩码、网关与 MAC。数值来自板端 API。板卡和 OS 标签则是代码中的声明，不单独作为 openvela 来源证明。')
block(.8, 2.65, 5.6, 'device_info', '内核 / 架构 / 运行时间\n堆总量、已用、空闲、最大空闲块\n结构化 JSON，便于核对和解释')
block(6.8, 2.65, 5.6, 'network_status', 'eth0 注册、UP、RUNNING 状态\nIPv4、网关、掩码、MAC\n缺失字段返回 null，而非虚构数据')
text(1.0, 5.4, 11.2, .85, '能力边界：接口“已配置”不代表外网可达；本工具不测 DNS、丢包或链路速率。\n演示时用 ifconfig / free 对照，证明回答有设备数据依据。', 17, MUTED)

new('现场演示：一次问答完成一次诊断', '05 / 演示主画面',
    '建议提问：请调用 device_info 和 network_status，概括设备内存与网络配置。', 28,
    '这里替换为真实录屏片段或截图。先执行 ifconfig 和 free，再发起工具问答，最后显示 nsh 提示符和第二次 free。若精简日志不显示工具调用，不要把自然语言回答当作调用证据，可补已有调试记录或另附工具链路证明。任何截图都要来自实机。')
placeholder(.8, 2.55, 7.9, 3.97, '02', '工具问答与 NSH 返回', '补图：提问、设备状态回答、nsh>。\n同时保留工具调用证据；严禁用模拟输出冒充。')
text(9.05, 2.8, 3.3, .4, '三个观察点', 22, TEAL, True)
text(9.05, 3.55, 3.3, 2.3, '01  数据能与命令核对\n\n02  回答后回到 NSH\n\n03  可继续执行下一条命令', 18)

new('不是泛泛“死机”：定位实际任务栈', '06 / 深度工程案例 A',
    '配置名相似并不代表作用于同一个任务；必须回到真实执行路径。', 25,
    'set_llm 曾多次卡住。GDB 显示 getumask 处访问地址 0x3e，任务信息受到破坏。追踪发现 builtin 在 NSH 当前任务运行，启动栈使用 INIT_STACKSIZE，而不是已配置为 32KB 的 SYSTEM_NSH_STACKSIZE。实际启动栈仍为 2KB。修正后 ps 观测峰值 3572 字节，超过旧栈容量，命令也能返回。这是栈不足的重要证据，不是对所有历史故障的一概归因。')
block(.8, 2.60, 6.05, '现象 → 证据 → 修正', 'set_llm 卡住；getumask 访问异常\nINIT 栈实际 2 KB，不是预期 32 KB\n调整真实启动栈，启用栈着色验证')
text(1.0, 5.30, 5.6, .85, '实机 ps：NSH 栈峰值 3572 B\n超过旧 2048 B 配置。', 21, TEAL, True)
placeholder(7.4, 2.55, 5.15, 3.97, '03', 'GDB 异常 + ps 栈证据', '补图：mepc / mtval 与修正后的 ps。\n可分为上下两张，保留关键数值。')

new('free 无限打印，根因在堆布局', '07 / 深度工程案例 B',
    '不在显示层截断输出，而是修正内存分配器初始化配置。', 25,
    'free 反复打印同一行。GDB 直接证实 procfs 的 Umem 节点 next 指向自身。检查板级堆分配发现，在当时配置下独立内核堆和用户堆落到同一片 SRAM。flat 模式改为共享分配器并全量重编译，用户确认 free 恢复。现场链表证据与板级代码一起说明了问题，而不是仅凭最后一条日志猜测。')
block(.8, 2.60, 6.0, '链表自环 → 重叠堆 → 共享分配器', 'GDB：g_procfs_meminfo->next 指向自己\n检查 heap / kheap 返回的 SRAM 区域\n关闭独立内核堆，完整重编译')
text(1.0, 5.32, 5.7, .8, '结果：实机反馈 free 恢复。\n避免以“断开链表”掩盖内存损坏。', 19, TEAL, True)
placeholder(7.4, 2.55, 5.15, 3.97, '04', '自环证据 + 正常 free', '补图：next == self 为 1 的 GDB 记录，\n以及修正后只输出一次并返回 NSH。')

new('跨层适配，让云端请求真正跑通', '08 / 集成工作量',
    '解决的不是单一 API，而是输入、存储、网络、安全传输和时间的一组约束。', 24,
    'USB 软件队列扩大并统一日志输出路径。TLS 启用实际服务端需要的 ECDHE 套件并清理重建库。NTP 提供真实日期，耗时使用单调时钟，避免校时导致超时判断跳变。SmartFS 的短文件名修复已编译，还要进行最终板端回归。TLS 握手成功不代表证书安全已经产品化。')
block(.8, 2.65, 5.7, 'USB 与文件系统', 'USB RX/TX 环形队列 64 → 2048 B\n统一正常日志与控制台输出路径\nSmartFS 短名称兼容与写入错误检查')
block(6.9, 2.65, 5.6, 'TLS 与时间', '启用 ECDHE，清理重建依赖归档\nNTP 校时，不再强设固定日期\nCLOCK_MONOTONIC 统计请求耗时')
text(1, 5.44, 11.3, .85, '已有实机握手：TLSv1.2 / TLS-ECDHE-RSA-WITH-AES-256-GCM-SHA384\n边界：证书校验仍为 VERIFY_OPTIONAL，尚未完成生产安全配置。', 15, MUTED)

new('一次 ask，一次可结束的工作单元', '09 / 面向无屏板卡的运行方式',
    '同步执行、复用核心状态、释放请求缓冲区；避免把命令行变成假交互界面。', 22,
    '原来问答后看见 vela 提示符，却不能继续输入。现在 CLI ask 走同步 ReAct 路径，首次初始化核心组件，后续复用状态，结束后释放本次请求缓冲并回到 NSH。不启动 CLI 不需要的 cron、heartbeat 和 WebSocket 后台服务。这里是运行方式的针对性改进，不宣称重新发明 Agent。')
for x, n, title, body in [(0.85, '01', '进入', 'NSH 发起 ask\n首次完成核心初始化'), (4.95, '02', '执行', '复用 ReAct 循环\n模型请求与工具执行'), (9.0, '03', '返回', '释放请求缓冲\n输出一次并返回 NSH')]:
    text(x, 2.68, 2, .72, n, 36, TEAL, True)
    text(x, 3.75, 3.4, .42, title, 23, INK, True)
    text(x, 4.46, 3.5, 1.15, body, 17)
text(.9, 6.08, 11.4, .45, '取舍：此模式不提供常驻提醒与 WebSocket 服务；验证重点是连续请求和资源稳定性。', 15, MUTED)

new('AI 协作不是替代验证，而是加快验证', '10 / 开发过程与可复用成果',
    '需求拆解 → 代码建议 → 构建 → 实机反馈 → GDB 证据 → 定点修正。', 20,
    '这个项目经过多轮失败和定位。AI 帮助跨仓检索、生成补丁与整理假设，但最终判断依据是实机和调试器。准备把经验沉淀为开发电脑上的 bringup-debug Skill，包含前置条件、操作步骤与验收标准。当前 Skill 和日志提交还在整理，不能把板端自带 Skill 当成自己的开发成果。')
block(.8, 2.62, 5.9, '已发生的协作工作', '跨仓配置与代码核对\n分析串口日志、异常寄存器和链表\n将修正编译后交由板端验证')
block(6.9, 2.62, 5.6, '待归集的开发 Skill', '版本与 ELF 一致性检查\n栈 / 堆 / USB / TLS 排障路径\n每一步的验收条件与失败分支')
text(1, 5.42, 11.2, .85, '工作量快照：3 个仓库、30 个已跟踪文件存在未提交改动，另有新配置与工具文件。\n这是当前工程覆盖范围，不是原创代码量，也不包括全部历史工作。', 16, MUTED)

new('用可核对的证据结束演示', '11 / 验证结果与当前边界',
    '展示已验证结果；把发布前补测项明确留在清单中。', 25,
    '已有用户实机日志证明 DHCP 获取地址，网关和公网域名各十次 ping 都收到响应，华为 MaaS TLS 和问答跑通，free 和日期恢复。这些是单次功能验证，不是压力测试。明天补录工具回答与命令输出对照、冷启动与连续请求。最新短文件名修复和构建配置一致性也需要验收。')
placeholder(.8, 2.55, 6.1, 3.97, '05', '网络 / 时间 / 连续请求', '补图：ifconfig、date、两次 ask、free。\n可用已有 ping 记录作辅证，不宣称长期零丢包。')
text(7.25, 2.70, 5.0, .4, '已有实机反馈', 21, TEAL, True)
text(7.25, 3.31, 5.0, 1.28, 'DHCP、局域网 / 公网 ping\nTLS 与问答、NSH 返回\nfree、日期恢复正常', 17)
text(7.25, 4.99, 5.0, .4, '最终补测', 21, RED, True)
text(7.25, 5.60, 5.0, .8, '新文件名、配置对齐、重启持久化\n连续请求资源变化、工具调用证据', 16)

new('交付的不只是启动画面', '12 / 总结',
    '把一个能跑的板卡，推进为能读取状态、能解释结果、能继续操作的 AI 终端。', 15,
    '项目的独特工作有三点：把设备真实状态接入 Agent，完成面向无屏板卡的同步问答链路，以及用异常和内存证据解决影响可靠性的移植问题。我们尊重上游成果，提交自己的增量代码、验证材料与可复用经验。下一步重点是发布验收和安全完善，不是堆叠未经验证的功能。')
block(1, 2.83, 3.5, '真实', '板端 API → JSON\n模型基于观测值解释')
block(5.0, 2.83, 3.5, '可控', 'ask 结束返回 NSH\n不依赖额外外设')
block(9.0, 2.83, 3.5, '可追溯', '故障证据对应修正\n代码与测试边界清楚')
text(1.15, 5.78, 11.0, .5, '项目仓库：zealsoftstudio / contest2026_124_shenyiyangdelaoshi', 17, TEAL, True)

new('附录：核心改动如何定位', 'APPENDIX A / 答辩备用，不计入视频',
    '以下为代码导航。最终提交需要补齐基线、跨仓 PR 与正式发布记录。', 0,
    '用于评委追问时定位代码，不纳入主视频。当前基线加工作区改动并非最终发布提交。详细路径与哈希见介绍文档。')
rows = [('目标配置 / 启动', 'nuttx: configs/usbnet/defconfig、esp32p4_bringup.c'), ('USB / 存储', 'nuttx: esp_usbserial.c、esp_board_spiflash.c'), ('DHCP 初始化', 'apps: netutils/netinit/Kconfig、netinit.c'), ('同步 Agent', 'ai_agent: agent_main.c、core/agent_loop.c'), ('设备工具', 'ai_agent: tools/tool_esp32p4.c、tool_registry.c'), ('时间 / 文件名', 'ai_agent: core/agent_trace.c、tools/skill_loader.c')]
for i, (label, body) in enumerate(rows):
    y = 2.6+i*.58
    text(.9, y, 2.5, .40, label, 16, TEAL, True)
    text(3.4, y, 9.0, .40, body, 15)

new('附录：发布前须如实说明的限制', 'APPENDIX B / 答辩备用，不计入视频',
    '把限制说清楚，比用未经验证的功能充数更有说服力。', 0,
    '当前依赖公网模型，不能离线推理；工具不能主动证明互联网可达；TLS 证书校验未产品化；PSRAM 没有作为用户堆。当前 .config 与 defconfig 的 SmartFS 名称长度不同，应对齐后复验，不要随意改变已格式化数据区。开发 Skill、日志和上游 PR 属于交付整理工作，尚不能声称完成。')
text(1, 2.65, 11.2, 3.5, '01  依赖云端模型和网络；天气、搜索密钥未配置。\n02  CLI ask 不运行后台 cron / heartbeat / WebSocket。\n03  TLS VERIFY_OPTIONAL 尚未完成生产级证书验证。\n04  PSRAM 未作为用户堆；不宣称已支持所有硬件外设。\n05  SmartFS：当前 .config 为 16，defconfig 为 48，发布前需对齐复验。\n06  新短文件名、连续请求与干净构建仍需最终验收。', 19)


# Structured prose is shared by the editable Word file and PDF.
sections = [
('作品概述', [
('作品名称', '基于 ESP32-P4 的 openvela AI 设备诊断助手。队伍：124 神一样的老师。材料日期：2026-09-16。定位：新硬件适配与 AI 应用结合的工程原型。'),
('核心目标', '面向没有屏幕、没有额外传感器的开发板，使用 USB 控制台发起自然语言请求，让 openvela Agent 调用板端设备工具，读取系统与网络配置，再由云端模型解释结果。完成一次问答后返回 NSH，继续执行系统命令。'),
('为什么值得做', '系统能启动、网口能联网，并不自动等于 Agent 可用。模型调用会同时触及任务栈、内存分配、存储、TLS、系统时间和终端生命周期。本项目针对这些交叉约束完成集成与修正，目标是提供一个可继续开发的板端 AI 诊断基线。'),
('差异化主张', '不以“首个 ESP32-P4 移植”作为宣传点，也不推测其他参赛作品的水平。差异化在于：设备状态结构化工具、同步 CLI Agent 链路，以及由 GDB 与实机反馈支撑的系统级问题定位。'),
('完成度声明', '已有实机反馈覆盖联网、云端问答、命令返回、free 与日期。最新短文件名修复已编译，尚待最终实机验收；完整工具演示证据、连续运行记录及干净构建需在提交前补齐。本材料是可补图的提交草稿，不是最终验收报告。'),
]),
('架构与使用场景', [
('数据链路', 'USB NSH → ai_agent ask → openvela Agent / ReAct → 华为 MaaS HTTPS 推理 → 工具调用 → 板端 API → JSON → 模型解释 → NSH。设备承担工具执行与请求编排，模型推理在云端，不宣称本地大模型推理。'),
('device_info', '使用 uname、mallinfo 和 CLOCK_MONOTONIC，提供内核版本、架构、运行时间、堆总量、已用、空闲及最大空闲块。board 和 os 字段为实现中的固定标签；证明 openvela 来源应依赖源码版本、构建依赖和实际使用的 Agent 能力，而不是这两个字符串。'),
('network_status', '通过 netlib 获取 eth0 注册状态、UP / RUNNING 标志、IPv4、网关、掩码与 MAC。查询失败的地址字段返回 null。它报告接口状态与配置，不主动测量外网连通性、DNS、速率或丢包，不能把 has_ipv4 等同于互联网可达。'),
('数据与接口', '配置和会话存于 /data/agent，底层为 SmartFS。已测试的推理入口为 api.modelarts-maas.com/v2/chat/completions，模型标识为 deepseek-v4-flash。模型名称与服务可用性以用户账号实际开通情况为准；密钥由演示者私下配置，不能录入视频或提交仓库。'),
('建议演示问题', '“请调用 device_info 查看内存和运行时间。”随后询问“请调用 network_status，说明当前 IP、网关与接口状态，不要把接口 UP 当成外网可达。”使用 ifconfig、free 对照；数值可能因请求期间的资源使用和采样时刻而变化，不应要求完全相等。'),
]),
('上游基础与本项目贡献', [
('尊重开源归属', 'ESP32-P4 架构、芯片驱动及 HAL 以现有 NuttX / Espressif 和社区实现为基础；Agent 主体、工具注册与 ReAct 框架来自 openvela；JSON、TLS、校时复用 cJSON、mbedTLS 和 NTP。不得把这些组件整体计为本项目原创。正式提交应补充参考 PR 链接、版本及许可说明。'),
('贡献 A：目标板集成', '面向 Function EV 板整理 USB Console、timer、EMAC、DHCP、procfs、SmartFS 与 Agent 的目标配置；改动涉及内核、应用网络初始化及 Agent 构建配置。工作目标是让板卡启动后的服务依赖顺序与实际可用性一致。'),
('贡献 B：真实状态工具', '新增 tool_esp32p4.c / .h，实现 device_info 与 network_status，并接入工具注册。工作不是仅添加提示词，而是把真实系统调用返回的数据转换为模型可消费的结构化结果。'),
('贡献 C：命令生命周期', '增加同步 agent_loop_ask 路径，由 CLI 直接完成 ReAct 请求，复用核心初始化状态并释放请求级缓冲区。避免一次问答启动不必要的常驻服务，回答结束交还 NSH。同步模式不提供后台 cron、heartbeat 或 WebSocket 服务。'),
('贡献 D：跨层可靠性', '围绕实际症状完成启动任务栈配置、共享堆布局、USB 队列与输出路径、TLS 套件、时间基准和文件名边界的分析与修正。各项修正的证据强弱与验证状态在后文分别注明，不把全部历史卡顿归为单一原因。'),
]),
('关键故障定位：任务栈与内存', [
('案例 A：set_llm 卡住', 'GDB 显示异常位置 getumask + 18，mtval = 0x3e。代码路径检查发现，CONFIG_NSH_BUILTIN_AS_COMMAND 下 builtin 使用当前 NSH 任务；启动入口 nsh_main 使用 CONFIG_INIT_STACKSIZE，而此前只扩大了 CONFIG_SYSTEM_NSH_STACKSIZE。旧启动栈为 2048 字节。'),
('修正与证据', '将 CONFIG_INIT_STACKSIZE 设为 32768，开启栈着色与调试符号。用户实机 ps 显示 NSH 栈可用 32704 字节、观测使用峰值 3572 字节，超过旧 2 KB 配置；set_llm 和 config_show 返回 NSH。结合执行路径、栈峰值与修正效果，支持启动任务栈不足这一判断。'),
('案例 B：free 无限打印', '用户 GDB 直接观测 g_procfs_meminfo = 0x4ff5dac0，节点 next 同样为 0x4ff5dac0，比较表达式为 1，说明 Umem 链表节点自环。该证据解释了 meminfo_read 不断输出相同行的现象。'),
('修正与证据', '在当时关闭 PSRAM 用户堆但启用独立内核堆的配置下，板级用户堆和内核堆返回重叠 SRAM。flat 模式关闭 CONFIG_MM_KERNEL_HEAP，改为共享分配器并完整清理重编译；ELF 检查不再保留独立内核堆符号，用户随后确认 free 正常。不是简单修改 free 或人为断开链表。'),
('截图位 A / B', '请补：GDB 异常寄存器、修正后的 ps；procfs 节点自环记录、修正后的 free 与 nsh>。旧记录仅用于解释定位过程，不作为当前镜像仍有同样故障的证据。'),
]),
('跨层集成：USB、TLS、时间与存储', [
('USB 控制台', '原软件环形队列为 64 字节，有效容量 63 字节；修改 RX/TX 缓冲为 2048 字节，并通过 SYSLOG_CONSOLE 统一正常日志输出路径。该修正针对长命令丢字和输出交错；最终仍应补连续长文本回显测试，不宣称绝对不丢字。'),
('TLS 实际服务互通', '启用 ECDHE-RSA / ECDHE-ECDSA，排查配置变化后静态库未充分重建的问题。用户日志已出现 TLSv1.2 / TLS-ECDHE-RSA-WITH-AES-256-GCM-SHA384 握手成功，并收到模型回答。握手成功不等于生产级安全完成；当前仍为 MBEDTLS_SSL_VERIFY_OPTIONAL。'),
('时间与耗时分离', '删除把系统时钟强设为 2026-02-28 的补偿逻辑；冷启动时间无效时启动 NTP 并短时等待，失败返回明确提示。LLM 耗时与 trace 改用 CLOCK_MONOTONIC，避免校时跳变造成几十亿毫秒的假超时。用户已反馈日期恢复至 2026-09-16。'),
('文件名边界', '当前已构建镜像的 SMARTFS_MAXNAMLEN 为 16；文件系统写名称时预留终止符，16 字节名称存在截断边界。内置名称改为 daily-brief.md、skill-create.md、sys-health.md，增加 errno 及写入/关闭错误检查。修正已编译，最终实机验证待补；不自动删除旧文件或重新格式化数据。'),
('发布前配置风险', '本次核对发现 usbnet/defconfig 的 SMARTFS_MAXNAMLEN 为 48，而当前 .config 为 16。应结合已有数据区格式对齐正式构建配置并复验，不能将两者视作已经一致。该项尚未在本次文档制作中修改。'),
]),
('工作量与验证矩阵', [
('工作量口径', '2026-09-16 工作区快照：nuttx 有 9 个、apps 有 3 个、packages/ai_agent 有 18 个已跟踪文件存在未提交改动，共 30 个；另有 usbnet 配置目录及 tool_esp32p4.c / .h 新文件。统计只反映当前增量范围，不代表全部移植代码，也不用于声称原创行数或人天。'),
('七类工程工作', '1. 板级启动与构建配置；2. USB 终端；3. timer / EMAC / DHCP；4. 可写存储；5. Agent CLI 与设备工具；6. TLS / 时间；7. 栈堆定位、回归与交付材料。主要难点是层间约束，不是单文件代码长度。'),
('已获得实机反馈', 'NSH 与联网；DHCP 获取 192.168.1.136；网关及公网域名各一次 10 包 ping 测试均 10 发 10 收；TLS / 云端问答；同步命令返回；free 与日期正常。IP 为当次 DHCP 结果，不是固定配置；零丢包只限上述短测试，不代表长期可靠性。'),
('必须补录 / 补测', '两项设备工具与命令输出的对照；最新短文件名及内容读取；重启后配置持久化（截图不含密钥）；连续 10 次 ask 前后 free / ps；无效输入与网络异常的可恢复性；干净目录构建和配置一致性。记录 pass / fail 与实际结果，不能预填通过。'),
('AI 开发协作', 'AI 用于需求拆解、跨仓检索、补丁生成、日志分析与文档整理；开发者负责板端实验和反馈，调试器提供现场证据。计划归集开发 Skill 与脱敏日志，当前未完成提交，不用板端自带天气 Skill 替代开发成果。'),
]),
('演示方案与已知限制', [
('视频顺序', '主 PPT 13 页，建议旁白约 4 分 47 秒，可根据实机片段调整到 5 分钟以内。先说明场景与原创边界，再展示架构、真实工具、两个深度定位案例和完成度。附录 2 页只供答辩，不计入主视频。'),
('五处实机图', '第 1 页板卡连接；第 6 页工具问答与 NSH 返回；第 7 页任务栈证据；第 8 页堆链表与 free；第 12 页联网、时间与连续请求。所有占位均明确标记为待补图，没有生成虚构实机结果。'),
('能力限制', '依赖公网与云端模型，不支持离线大模型推理。天气、搜索等服务密钥未配置；注册表中存在的工具不等于本板具备相应硬件。PSRAM 当前未作为用户堆，不宣称 32 MB 可用 Agent 堆。没有完成全外设验证、长稳测试或性能基准。'),
('安全限制', 'TLS 证书验证尚未产品化，API Key 存储保护也不应视为完成。此前开发交流中出现过密钥，应撤销/轮换，并检查文档、日志、录屏、配置文件及 Git 历史。只演示脱敏信息，不发布含用户数据的 SmartFS 分区。'),
('提交清单', '正式 README；可复现源码与固定版本；跨仓 PR 及许可归属；固件和校验值；开发 Skill 与脱敏会话日志；最终测试记录；介绍文档、PPT 和不超过 5 分钟的视频。本次仅制作介绍材料，不表示已经提交或已满足全部验收项。'),
]),
('代码证据与版本记录', [
('关键代码路径', 'nuttx/boards/risc-v/esp32p4/esp32p4-function-ev-board/configs/usbnet/defconfig；同板 src/esp32p4_bringup.c；nuttx/arch/risc-v/src/esp32p4/espressif/esp_usbserial.c；nuttx/boards/risc-v/esp32p4/common/src/esp_board_spiflash.c。'),
('Agent 证据路径', 'packages/ai_agent/src/agent_main.c；src/core/agent_loop.c 与 agent_trace.c；src/tools/tool_esp32p4.c 与 tool_registry.c；src/tools/skill_loader.c；src/infra/vela_tls.c。DHCP 相关增量位于 apps/netutils/netinit。'),
('基线快照（不是最终发布版本）', 'nuttx: 10c4e090243；apps: dcc6a95c3b323e533c98fde8fb209f99e24f0fdd；packages/ai_agent: e65550f18759f086d7f544edcf17d1e31223244f。三仓均叠加本地修改，单独检出这些提交不能复现本项目。'),
('当前待最终验收镜像', 'nuttx/nuttx.bin，695176 字节；SHA-256: d8b6cdd16e46cc0a0cef14e6e94a5702bf6da6c4cada3000e93e6b8eeefabae7。构建记录为 2026-09-16 21:58:33 +0800。最终配置对齐或重新构建后必须更新此记录和截图版本。'),
('参考与归属', '上游：github.com/open-vela/nuttx、github.com/open-vela/packages_ai_agent。比赛总览与提交指南位于 github.com/open-vela/docs 的 dev-ai-contest-2026 分支 zh-cn/contest_2026/。本材料依据本地源码、构建记录及用户实机反馈整理；未验证的同行比较与首次适配主张均不采用。'),
]),
]


def wrap(value, width, size, bold=False):
    font = ImageFont.truetype(str(BOLD if bold else FONT), round(size*2))
    lines = []
    for paragraph in value.split('\n'):
        line = ''
        for char in paragraph:
            if line and font.getlength(line+char) > width*144:
                lines.append(line)
                line = char
            else:
                line += char
        lines.append(line)
    return lines


def build_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    pdfmetrics.registerFont(TTFont('Chinese', str(PDF_FONT)))
    cv = canvas.Canvas(str(OUT/'presentation-preview.pdf'), pagesize=(W*72, H*72))
    preview_dir = OUT/'previews'
    preview_dir.mkdir(exist_ok=True)
    previews = []
    for idx, spec in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        im = Image.new('RGB', (1600, 900), '#'+BG)
        draw = ImageDraw.Draw(im)
        for item in spec['items']:
            x, y, w, h = (item[k] for k in ('x', 'y', 'w', 'h'))
            assert x >= 0 and y >= 0 and x+w <= W+.01 and y+h <= H+.01
            if item['kind'] == 'box':
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(item['fill'])
                if item['stroke']:
                    shape.line.color.rgb = RGBColor.from_string(item['stroke'])
                    shape.line.width = Pt(.7)
                else:
                    shape.line.fill.background()
                draw.rectangle((x*120, y*120, (x+w)*120, (y+h)*120), fill='#'+item['fill'], outline='#'+item['stroke'] if item['stroke'] else None)
                cv.setFillColor('#'+item['fill'])
                cv.setStrokeColor('#'+(item['stroke'] or item['fill']))
                cv.rect(x*72, (H-y-h)*72, w*72, h*72, fill=1, stroke=int(bool(item['stroke'])))
            else:
                lines = wrap(item['value'], w, item['size'], item['bold'])
                leading = item['size']*1.32
                assert len(lines)*leading <= h*72+4, (idx, item['value'], lines, h)
                shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
                tf = shape.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                for j, line in enumerate(lines):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    p.text = line
                    p.line_spacing = Pt(leading)
                    p.space_before = p.space_after = Pt(0)
                    for run in p.runs:
                        run.font.name = 'Microsoft YaHei'
                        run.font.size = Pt(item['size'])
                        run.font.bold = item['bold']
                        run.font.color.rgb = RGBColor.from_string(item['color'])
                        ea = OxmlElement('a:ea')
                        ea.set('typeface', 'Microsoft YaHei')
                        run._r.get_or_add_rPr().append(ea)
                    font = ImageFont.truetype(str(BOLD if item['bold'] else FONT), round(item['size']*120/72))
                    draw.text((x*120, y*120+j*leading*120/72), line, fill='#'+item['color'], font=font)
                    cv.setFont('Chinese', item['size'])
                    cv.setFillColor('#'+item['color'])
                    cv.drawString(x*72, (H-y)*72-item['size']-j*leading, line)
                shape.name = ('Screenshot placeholder text: ' if '截图位' in item['value'] else 'Text: ')+item['value'][:45]
        slide.notes_slide.notes_text_frame.text = f"建议时长：{spec['seconds']} 秒（0 为答辩附录）\n{spec['notes']}"
        im.save(preview_dir/f'{idx:02d}.png')
        previews.append(im)
        cv.showPage()
    prs.save(str(OUT/'ESP32P4-openvela-AI-diagnostics.pptx'))
    cv.save()
    contact = Image.new('RGB', (1200, 225*5), '#D9E2E5')
    for i, im in enumerate(previews):
        im.thumbnail((400, 225))
        contact.paste(im, ((i%3)*400, (i//3)*225))
    contact.save(OUT/'slide-overview.png')


def build_docs():
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Microsoft YaHei', DPt(10.5)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = DPt(7)
    normal.paragraph_format.line_spacing = 1.12
    for style in ['Title', 'Heading 1', 'Heading 2']:
        doc.styles[style].font.name = 'Microsoft YaHei'
        doc.styles[style]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        doc.styles[style].font.color.rgb = DColor.from_string(TEAL)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = DInches(.65)
    sec.left_margin = sec.right_margin = DInches(.75)
    sec.header.paragraphs[0].text = '队伍 124 神一样的老师 | openvela AI 设备诊断助手 | 提交前补图稿'
    footer = sec.footer.paragraphs[0]
    footer.text = '2026-09-16  ·  '
    fld = DX('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    pdf_style = ParagraphStyle('body', fontName='Chinese', fontSize=10.5, leading=16, spaceAfter=10, wordWrap='CJK')
    heading_style = ParagraphStyle('h1', parent=pdf_style, fontSize=23, leading=30, textColor='#'+TEAL, spaceAfter=18)
    sub_style = ParagraphStyle('h2', parent=pdf_style, fontSize=12, leading=18, textColor='#'+TEAL, spaceAfter=5)
    story = []
    markdown = ['# 基于 ESP32-P4 的 openvela AI 设备诊断助手', '', '> 2026-09-16 提交前补图稿。真实结果、代码检查和待测项目分开标注。', '']
    for idx, (title, paragraphs) in enumerate(sections):
        if idx:
            doc.add_page_break()
            story.append(PageBreak())
        doc.add_heading(f'{idx+1:02d}  {title}', level=1)
        story.append(Paragraph(f'{idx+1:02d}  {title}', heading_style))
        markdown.extend([f'## {idx+1}. {title}', ''])
        for label, content in paragraphs:
            doc.add_heading(label, level=2)
            doc.add_paragraph(content)
            story.extend([Paragraph(escape(label), sub_style), Paragraph(escape(content), pdf_style)])
            markdown.extend([f'### {label}', content, ''])
    doc.save(str(OUT/'project-introduction.docx'))
    def pdf_footer(c, d):
        c.setFont('Chinese', 9)
        c.setFillColor('#'+MUTED)
        c.drawString(45, 24, '队伍 124 | openvela AI 设备诊断助手 | 提交前补图稿')
        c.drawRightString(A4[0]-45, 24, str(d.page))
    SimpleDocTemplate(str(OUT/'project-introduction.pdf'), pagesize=A4, topMargin=42, bottomMargin=44, leftMargin=45, rightMargin=45).build(story, onFirstPage=pdf_footer, onLaterPages=pdf_footer)
    (OUT/'project-introduction.md').write_text('\n'.join(markdown), encoding='utf-8')
    script = ['# 视频旁白与补图清单', '', '主视频使用第 1-13 页；第 14-15 页为答辩附录。', f'建议时长合计 {sum(s["seconds"] for s in slides)} 秒，需结合实机片段控制在 5 分钟内。', '', '## 补图操作', '', 'PPT 所有文字与框均可编辑。选择灰色截图框及其中提示文字删除，再插入对应实机图片；按原框裁剪，不遮挡页标题和页脚。保留原始录屏，日志及截图须脱敏。', '', '截图编号：01 板卡连接；02 工具问答；03 任务栈；04 堆链表；05 回归结果。不要用示意图或旧故障画面替代当前成功证据。', '']
    for i, s in enumerate(slides, 1):
        script.extend([f'## 第 {i} 页：{s["title"]}（{s["seconds"]} 秒）', s['notes'], ''])
    script.extend(['## 发布前检查', '', '- 填补实机截图，删除截图位提示文字。', '- 实际完成 Skill / 日志 / 回归后再更新相应状态，未完成的继续标为待完成。', '- 对齐 SMARTFS_MAXNAMLEN 与数据区格式后重建、复验，更新镜像校验值。', '- PPT 和 Word 人工改动后，从 PowerPoint / Word 重新导出 PDF，预览 PDF 不会自动跟随修改。', '- 核对队伍中文名与报名信息；旁白按个人习惯改为“我”或“我们”。', '- 最终以 PowerPoint 打开检查字体和换行；本目录 PNG/PDF 是同源布局预览，不是 Office 原生渲染。'])
    (OUT/'video-script-and-screenshots.md').write_text('\n'.join(script), encoding='utf-8')


if __name__ == '__main__':
    build_deck()
    build_docs()
    (OUT/'layout-data.json').write_text(json.dumps(slides, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Generated {len(slides)} slides, {len(sections)} document sections in {OUT}')
